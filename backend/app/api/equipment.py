import asyncio
import io
import ipaddress
import re
import unicodedata
from collections import Counter
from datetime import date, datetime, timedelta
from zipfile import BadZipFile

import pytesseract
import qrcode
from PIL import Image, ImageOps, UnidentifiedImageError
from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, Field, ValidationError
from pymongo.errors import DuplicateKeyError

from app.ai.engine import equipment_row, get_engine, load_ticket_rows
from app.ai.inventory_normalizer import normalize_inventory_fields
from app.api.deps import parse_id, require_admin, require_staff
from app.core.config import get_settings
from app.core.timeutil import utcnow
from app.db import next_sequence
from app.models import Device, Equipment, MunicipalUser, Office, StaffUser, Ticket
from app.models.enums import EquipmentStatus, EquipmentType
from app.schemas.admin import EquipmentIn, EquipmentOut
from app.schemas.common import Message
from app.services import audit
from app.services.serializers import equipment_out, ticket_out
from app.schemas.ticket import TicketOut

router = APIRouter(prefix="/equipment", tags=["equipos"])


class EquipmentDetailOut(BaseModel):
    equipment: EquipmentOut
    risk: float | None
    risk_factors: list[str]
    tickets: list[TicketOut]


class EquipmentImportError(BaseModel):
    row: int
    patrimonial_code: str | None = None
    message: str


class EquipmentNormalizationExample(BaseModel):
    row: int
    field: str
    original: str
    normalized: str
    method: str


class EquipmentImportResult(BaseModel):
    processed: int
    imported: int
    rejected: int
    normalized: int = 0
    normalizations: list[EquipmentNormalizationExample] = Field(default_factory=list)
    more_normalizations: int = 0
    errors: list[EquipmentImportError]
    more_errors: int = 0


class EquipmentImportOfficeRef(BaseModel):
    office_code: str
    office_name: str
    zone_id: str | None
    zone_name: str | None
    head_name: str | None
    import_enabled: bool


class EquipmentOcrOut(BaseModel):
    detected_code: str | None
    matched: bool
    equipment_id: str | None = None
    candidates: list[str] = []


class RetirementHistoryItem(BaseModel):
    number: str
    created_at: datetime
    subject: str
    status: str
    resolution_type: str | None = None
    resolution_notes: str | None = None


class EquipmentRetirementReport(BaseModel):
    generated_at: datetime
    equipment: EquipmentOut
    total_incidents: int
    resolved_incidents: int
    incidents_365d: int
    incidents_90d: int
    resolution_counts: dict[str, int]
    risk_30d: float | None
    risk_factors: list[str]
    indicators: list[str]
    recommendation: str
    technical_conclusion: str
    disclaimer: str
    history: list[RetirementHistoryItem]


_HEADER_ALIASES = {
    "zona": "zona",
    "nombre_zona": "zona",
    "zona_id": "zona_id",
    "id_zona": "zona_id",
    "oficina": "oficina",
    "oficina_codigo": "oficina_codigo",
    "codigo_oficina": "oficina_codigo",
    "codigo_patrimonial": "codigo_patrimonial",
    "cod_patrimonial": "codigo_patrimonial",
    "patrimonial_code": "codigo_patrimonial",
    "tipo": "tipo",
    "tipo_equipo": "tipo",
    "area": "area",
    "dispositivo": "dispositivo",
    "marca": "marca",
    "modelo": "modelo",
    "tamano_pantalla": "tamano_pantalla",
    "tamaño_pantalla": "tamano_pantalla",
    "nombre_equipo": "nombre_equipo",
    "ip": "ip",
    "direccion_ip": "ip",
    "hostname": "hostname",
    "mac": "mac",
    "direccion_mac": "mac",
    "cpu": "cpu",
    "procesador": "cpu",
    "ram_gb": "ram_gb",
    "ram": "ram_gb",
    "memoria_ram_gb": "ram_gb",
    "almacenamiento_gb": "almacenamiento_gb",
    "almacenamiento": "almacenamiento_gb",
    "disco_gb": "almacenamiento_gb",
    "sistema_operativo": "sistema_operativo",
    "so": "sistema_operativo",
    "fecha_adquisicion": "fecha_adquisicion",
    "garantia_hasta": "garantia_hasta",
    "estado": "estado",
    "propiedad": "propiedad",
    "responsable": "responsable",
    "usuario": "responsable",
    "usuario_responsable": "responsable",
    "codigo_responsable": "responsable_codigo",
    "responsable_codigo": "responsable_codigo",
    "responsable_tipo": "responsable_tipo",
    "tipo_responsable": "responsable_tipo",
    "criticidad": "criticidad",
    "notas": "notas",
}

_SUPPORTED_EQUIPMENT_TYPES = {
    EquipmentType.CPU,
    EquipmentType.MONITOR,
    EquipmentType.MOUSE,
    EquipmentType.TECLADO,
    EquipmentType.IMPRESORA,
    EquipmentType.LAPTOP,
}

_EQUIPMENT_TYPE_ALIASES = {
    "CPU": EquipmentType.CPU,
    "COMPUTADORA": EquipmentType.CPU,
    "COMPUTADOR": EquipmentType.CPU,
    "DESKTOP": EquipmentType.CPU,
    "PC": EquipmentType.CPU,
    "LAPTOP": EquipmentType.LAPTOP,
    "PORTATIL": EquipmentType.LAPTOP,
    "IMPRESORA": EquipmentType.IMPRESORA,
    "MONITOR": EquipmentType.MONITOR,
    "MOUSE": EquipmentType.MOUSE,
    "RATON": EquipmentType.MOUSE,
    "TECLADO": EquipmentType.TECLADO,
}

_STATUS_ALIASES = {
    "OPERATIVO": EquipmentStatus.OPERATIVO,
    "EN_REPARACION": EquipmentStatus.EN_REPARACION,
    "EN REPARACION": EquipmentStatus.EN_REPARACION,
    "BAJA": EquipmentStatus.BAJA,
}


def _normal(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").strip())
    return "".join(ch for ch in text if not unicodedata.combining(ch)).lower().replace(" ", "_")


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _ocr_normalize(value: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", (value or "").upper())


def _ocr_candidates(text: str) -> list[str]:
    lines = [line.strip().upper() for line in text.splitlines() if line.strip()]
    prioritized = [line for line in lines if "PATRIMON" in line or "CODIGO" in line or "CÓDIGO" in line]
    pool = prioritized + lines
    found: list[str] = []
    for line in pool:
        for token in re.findall(r"[A-Z0-9][A-Z0-9._/-]{3,39}", line):
            cleaned = token.strip("._/-")
            compact = _ocr_normalize(cleaned)
            if 4 <= len(compact) <= 40 and compact not in {"PATRIMONIAL", "CODIGO", "MUNICIPALIDAD", "PROVINCIAL", "CASMA"}:
                if cleaned not in found:
                    found.append(cleaned)
    return found[:12]


def _prepare_ocr_image(raw: bytes) -> Image.Image:
    try:
        image = Image.open(io.BytesIO(raw))
        image.load()
    except (UnidentifiedImageError, OSError, ValueError):
        raise HTTPException(status_code=422, detail="No se pudo leer la imagen.")

    if image.width < 40 or image.height < 40:
        raise HTTPException(status_code=422, detail="La imagen es demasiado pequeña para leer la etiqueta.")

    image = image.convert("RGB")
    max_side = 3200
    if max(image.size) > max_side:
        image.thumbnail((max_side, max_side))

    gray = ImageOps.grayscale(image)
    gray = ImageOps.autocontrast(gray)
    if max(gray.size) < 1800:
        gray = gray.resize((gray.width * 2, gray.height * 2))
    return gray


def _parse_excel_date(value, field: str) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = _cell_text(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    raise ValueError(f"{field}: use AAAA-MM-DD o DD/MM/AAAA.")


def _optional_float(value, field: str) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        raise ValueError(f"{field}: debe ser numérico.")


def _equipment_type(value) -> EquipmentType:
    key = _normal(value).upper()
    if key in _EQUIPMENT_TYPE_ALIASES:
        return _EQUIPMENT_TYPE_ALIASES[key]
    try:
        equipment_type = EquipmentType(key)
    except ValueError:
        equipment_type = None
    if equipment_type not in _SUPPORTED_EQUIPMENT_TYPES:
        allowed = ", ".join(t.value for t in sorted(_SUPPORTED_EQUIPMENT_TYPES, key=lambda item: item.value))
        raise ValueError(f"Tipo de equipo no válido para el Área TI. Valores permitidos: {allowed}.")
    return equipment_type


def _ensure_supported_type(equipment_type: EquipmentType) -> None:
    if equipment_type not in _SUPPORTED_EQUIPMENT_TYPES:
        allowed = ", ".join(t.value for t in sorted(_SUPPORTED_EQUIPMENT_TYPES, key=lambda item: item.value))
        raise HTTPException(status_code=422, detail=f"Tipo de equipo no permitido. Use: {allowed}.")


async def _next_inventory_id() -> str:
    seq = await next_sequence("equipment-ti")
    return f"TI-{seq:06d}"


def _equipment_status(value) -> EquipmentStatus:
    if value in (None, ""):
        return EquipmentStatus.OPERATIVO
    key = _normal(value).upper()
    if key in _STATUS_ALIASES:
        return _STATUS_ALIASES[key]
    try:
        return EquipmentStatus(key)
    except ValueError:
        raise ValueError("Estado no válido. Use OPERATIVO, EN_REPARACION o BAJA.")


async def _get(equipment_id: str) -> Equipment:
    eid = parse_id(equipment_id)
    eq = await Equipment.get(eid) if eid else None
    if not eq:
        raise HTTPException(status_code=404, detail="Equipo no encontrado.")
    if not eq.inventory_id:
        eq.inventory_id = await _next_inventory_id()
        await eq.save()
    return eq


async def _validate_office(office_id: str | None) -> Office:
    if not office_id:
        raise HTTPException(status_code=422, detail="Debe seleccionar una oficina.")
    oid = parse_id(office_id)
    office = await Office.get(oid) if oid else None
    if not office:
        raise HTTPException(status_code=422, detail="Oficina no válida.")
    return office


async def _assignment_data(data: dict, office: Office) -> dict:
    result = dict(data)
    result["office_id"] = office.id
    result["area"] = (result.get("area") or office.name).strip()
    result["device_label"] = (result.get("device_label") or result.get("type").value).strip()
    result["property_type"] = (result.get("property_type") or "MUNICIPALIDAD PROVINCIAL DE CASMA").strip()

    responsible_type = str(result.get("responsible_type") or "USUARIO").upper().strip()
    responsible_name = (result.get("responsible_name") or "").strip()
    responsible_id_raw = result.get("responsable_id")

    if responsible_type == "OFICINA":
        responsible_name = responsible_name or office.name
        result["responsable_id"] = None
    elif responsible_type == "JEFE":
        responsible_name = responsible_name or office.head_name or office.name
        result["responsable_id"] = None
    elif responsible_type == "USUARIO":
        if responsible_id_raw:
            uid = parse_id(str(responsible_id_raw))
            municipal_user = await MunicipalUser.get(uid) if uid else None
            if not municipal_user or not municipal_user.active:
                raise HTTPException(status_code=422, detail="El usuario responsable no existe o está inactivo.")
            if municipal_user.office_id != office.id:
                raise HTTPException(status_code=422, detail="El usuario responsable no pertenece a la oficina seleccionada.")
            result["responsable_id"] = municipal_user.id
            responsible_name = municipal_user.full_name
        elif not responsible_name:
            raise HTTPException(status_code=422, detail="Seleccione un usuario responsable de la oficina.")
        else:
            # Compatibilidad temporal para registros antiguos todavía no vinculados.
            result["responsable_id"] = None
    else:
        raise HTTPException(status_code=422, detail="Tipo de responsable no válido.")

    result["responsible_type"] = responsible_type
    result["responsible_name"] = responsible_name
    return result


@router.get("", response_model=list[EquipmentOut])
async def list_equipment(
    office_id: str | None = None, type: EquipmentType | None = None, q: str | None = Query(None, max_length=60),
    _: StaffUser = Depends(require_staff),
):
    import re

    query: dict = {}
    if office_id and (oid := parse_id(office_id)):
        query["office_id"] = oid
    if type:
        query["type"] = type.value
    if q and q.strip():
        rx = {"$regex": re.escape(q.strip()), "$options": "i"}
        query["$or"] = [
            {"inventory_id": rx},
            {"patrimonial_code": rx},
            {"mac_address": rx},
            {"ip_address": rx},
            {"responsible_name": rx},
            {"hostname": rx},
            {"device_label": rx},
            {"area": rx},
            {"brand": rx},
            {"model": rx},
        ]
    offices = {o.id: o for o in await Office.find_all().to_list()}
    items = await Equipment.find(query).sort("patrimonial_code").limit(1000).to_list()
    for equipment in items:
        if not equipment.inventory_id:
            equipment.inventory_id = await _next_inventory_id()
            await equipment.save()
    return [
        equipment_out(
            e,
            offices[e.office_id].name if e.office_id in offices else None,
            str(offices[e.office_id].zone_id) if e.office_id in offices and offices[e.office_id].zone_id else None,
            offices[e.office_id].zone_name if e.office_id in offices else None,
        )
        for e in items
    ]


@router.post("", response_model=EquipmentOut, status_code=201)
async def create_equipment(request: Request, body: EquipmentIn, user: StaffUser = Depends(require_staff)):
    _ensure_supported_type(body.type)
    office = await _validate_office(body.office_id)
    data = await _assignment_data(body.model_dump(), office)
    data["inventory_id"] = await _next_inventory_id()
    eq = Equipment(**data)
    try:
        await eq.insert()
    except DuplicateKeyError:
        raise HTTPException(status_code=409, detail="Ya existe un equipo con ese código patrimonial.")
    await audit.record(request, "staff", "equipment.created", actor_id=str(user.id), actor_name=user.full_name, target_type="equipment", target_id=str(eq.id))
    return equipment_out(
        eq,
        office.name,
        str(office.zone_id) if office.zone_id else None,
        office.zone_name,
    )


@router.post("/ocr-patrimonial", response_model=EquipmentOcrOut)
async def ocr_patrimonial(
    file: UploadFile = File(...),
    _: StaffUser = Depends(require_staff),
):
    content_type = (file.content_type or "").lower()
    if not content_type.startswith("image/"):
        raise HTTPException(status_code=422, detail="Debe enviar una imagen de la etiqueta patrimonial.")

    max_bytes = get_settings().max_upload_mb * 1024 * 1024
    raw = await file.read(max_bytes + 1)
    if len(raw) > max_bytes:
        raise HTTPException(
            status_code=413,
            detail=f"La imagen no puede superar {get_settings().max_upload_mb} MB.",
        )
    if not raw:
        raise HTTPException(status_code=422, detail="La imagen está vacía.")

    image = _prepare_ocr_image(raw)
    try:
        text = await asyncio.to_thread(
            pytesseract.image_to_string,
            image,
            lang="spa+eng",
            config="--psm 6",
        )
    except pytesseract.pytesseract.TesseractNotFoundError:
        raise HTTPException(status_code=503, detail="El motor OCR no está disponible en el servidor.")
    except pytesseract.pytesseract.TesseractError:
        raise HTTPException(status_code=422, detail="No se pudo reconocer texto en la etiqueta.")

    text = (text or "").strip()
    candidates = _ocr_candidates(text)
    compact_text = _ocr_normalize(text)

    equipment_docs = await Equipment.get_pymongo_collection().find(
        {},
        {"patrimonial_code": 1},
    ).to_list(None)

    matched_code = None
    matched_id = None
    for doc in equipment_docs:
        code = str(doc.get("patrimonial_code") or "").strip().upper()
        normalized = _ocr_normalize(code)
        if normalized and normalized in compact_text:
            matched_code = code
            matched_id = str(doc["_id"])
            break

    if not matched_code:
        candidate_norms = [(_ocr_normalize(candidate), candidate) for candidate in candidates]
        for doc in equipment_docs:
            code = str(doc.get("patrimonial_code") or "").strip().upper()
            normalized = _ocr_normalize(code)
            if normalized and any(normalized == norm for norm, _ in candidate_norms):
                matched_code = code
                matched_id = str(doc["_id"])
                break

    detected = matched_code or (candidates[0].upper() if candidates else None)
    return EquipmentOcrOut(
        detected_code=detected,
        matched=bool(matched_code),
        equipment_id=matched_id,
        candidates=[candidate.upper() for candidate in candidates],
    )


@router.get("/export-xlsx")
async def export_equipment_xlsx(
    request: Request,
    zone_name: str | None = Query(None, max_length=120),
    office_id: str | None = None,
    area: str | None = Query(None, max_length=120),
    type: EquipmentType | None = None,
    status: EquipmentStatus | None = None,
    q: str | None = Query(None, max_length=60),
    admin: StaffUser = Depends(require_admin),
):
    query: dict = {}

    offices = await Office.find_all().to_list()
    office_map = {office.id: office for office in offices}

    if office_id:
        oid = parse_id(office_id)
        if not oid:
            raise HTTPException(status_code=422, detail="Oficina no válida.")
        query["office_id"] = oid
    elif zone_name and zone_name.strip():
        normalized_zone = _normal(zone_name)
        zone_office_ids = [
            office.id
            for office in offices
            if office.zone_name and _normal(office.zone_name) == normalized_zone
        ]
        query["office_id"] = {"$in": zone_office_ids}

    if area and area.strip():
        query["area"] = area.strip()
    if type:
        _ensure_supported_type(type)
        query["type"] = type.value
    if status:
        query["status"] = status.value
    if q and q.strip():
        rx = {"$regex": re.escape(q.strip()), "$options": "i"}
        query["$or"] = [
            {"inventory_id": rx},
            {"patrimonial_code": rx},
            {"mac_address": rx},
            {"ip_address": rx},
            {"responsible_name": rx},
            {"hostname": rx},
            {"device_label": rx},
            {"area": rx},
            {"brand": rx},
            {"model": rx},
        ]

    equipment = await Equipment.find(query).sort("patrimonial_code").to_list()
    user_ids = list({item.responsable_id for item in equipment if item.responsable_id})
    users = {
        user.id: user
        for user in await MunicipalUser.find({"_id": {"$in": user_ids}}).to_list()
    } if user_ids else {}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Margesi TI"

    headers = [
        "Zona", "Zona ID", "Oficina", "Código Oficina", "ID TI", "Código Patrimonial",
        "Tipo", "Área", "Dispositivo", "Marca", "Modelo", "Tamaño Pantalla",
        "Nombre Equipo", "Procesador", "Memoria RAM GB", "Almacenamiento GB",
        "Dirección IP", "Dirección MAC", "Sistema Operativo", "Propiedad",
        "Responsable", "Código Responsable", "Tipo Responsable",
        "Fecha Adquisición", "Garantía Hasta", "Estado", "Criticidad", "Notas",
    ]
    sheet.append(headers)

    header_fill = PatternFill("solid", fgColor="13233B")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for item in equipment:
        office = office_map.get(item.office_id)
        responsible = users.get(item.responsable_id) if item.responsable_id else None
        sheet.append([
            office.zone_name if office else None,
            str(office.zone_id) if office and office.zone_id else None,
            office.name if office else None,
            office.code if office else None,
            item.inventory_id,
            item.patrimonial_code,
            item.type.value,
            item.area,
            item.device_label,
            item.brand,
            item.model,
            item.specs.screen_size_inches,
            item.hostname,
            item.specs.cpu,
            item.specs.ram_gb,
            item.specs.storage_gb,
            item.ip_address,
            item.mac_address,
            item.specs.os,
            item.property_type,
            item.responsible_name,
            responsible.employee_code if responsible else None,
            item.responsible_type,
            item.acquired_on.date().isoformat() if item.acquired_on else None,
            item.warranty_until.date().isoformat() if item.warranty_until else None,
            item.status.value,
            item.criticality,
            item.notes,
        ])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {
        1: 24, 2: 26, 3: 32, 4: 18, 5: 14, 6: 22, 7: 15, 8: 24, 9: 22,
        10: 18, 11: 22, 12: 16, 13: 24, 14: 28, 15: 16, 16: 20, 17: 18,
        18: 20, 19: 22, 20: 34, 21: 30, 22: 20, 23: 18, 24: 18, 25: 18,
        26: 18, 27: 12, 28: 42,
    }
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()

    await audit.record(
        request,
        "staff",
        "equipment.exported_xlsx",
        actor_id=str(admin.id),
        actor_name=admin.full_name,
        target_type="equipment",
        target_id="bulk",
        exported=len(equipment),
        zone_name=zone_name,
        office_id=office_id,
        area=area,
        equipment_type=type.value if type else None,
        status=status.value if status else None,
        query=q,
    )

    filename = f"margesi-ti-casma-{datetime.now().strftime('%Y%m%d-%H%M')}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/import-references", response_model=list[EquipmentImportOfficeRef])
async def import_references(_: StaffUser = Depends(require_admin)):
    offices = await Office.find({"active": True}).sort("name").to_list()
    return [
        EquipmentImportOfficeRef(
            office_code=o.code,
            office_name=o.name,
            zone_id=str(o.zone_id) if o.zone_id else None,
            zone_name=o.zone_name,
            head_name=o.head_name,
            import_enabled=bool(o.zone_id or o.zone_name),
        )
        for o in offices
    ]


@router.post("/import-xlsx", response_model=EquipmentImportResult)
async def import_equipment_xlsx(
    request: Request,
    file: UploadFile = File(...),
    admin: StaffUser = Depends(require_admin),
):
    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(status_code=422, detail="El archivo debe tener extensión .xlsx.")

    raw = await file.read(10 * 1024 * 1024 + 1)
    if len(raw) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="El archivo Excel no puede superar 10 MB.")
    if not raw:
        raise HTTPException(status_code=422, detail="El archivo Excel está vacío.")

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError, ValueError):
        raise HTTPException(status_code=422, detail="No se pudo leer el archivo .xlsx.")

    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    first = next(rows, None)
    if not first:
        workbook.close()
        raise HTTPException(status_code=422, detail="El archivo no contiene filas.")

    headers: dict[str, int] = {}
    for index, value in enumerate(first):
        normalized = _normal(value)
        canonical = _HEADER_ALIASES.get(normalized)
        if canonical and canonical not in headers:
            headers[canonical] = index

    required = {"codigo_patrimonial", "tipo"}
    missing = sorted(required - set(headers))
    missing_office = "oficina" not in headers and "oficina_codigo" not in headers
    if missing or missing_office or ("zona_id" not in headers and "zona" not in headers):
        workbook.close()
        missing_text = ", ".join(missing) if missing else ""
        office_text = "oficina o codigo_oficina" if missing_office else ""
        zone_text = "zona o zona_id" if "zona_id" not in headers and "zona" not in headers else ""
        detail = ", ".join(x for x in (missing_text, office_text, zone_text) if x)
        raise HTTPException(status_code=422, detail="Faltan columnas obligatorias: " + detail + ".")

    offices = await Office.find({"active": True}).to_list()
    by_code = {_normal(o.code): o for o in offices}
    by_name: dict[str, list[Office]] = {}
    for office in offices:
        by_name.setdefault(_normal(office.name), []).append(office)

    municipal_users = await MunicipalUser.find({"active": True}).to_list()
    users_by_office_code: dict[tuple[str, str], MunicipalUser] = {}
    users_by_office_name: dict[tuple[str, str], list[MunicipalUser]] = {}
    for municipal_user in municipal_users:
        office_key = str(municipal_user.office_id)
        if municipal_user.employee_code:
            users_by_office_code[(office_key, _normal(municipal_user.employee_code))] = municipal_user
        users_by_office_name.setdefault((office_key, _normal(municipal_user.full_name)), []).append(municipal_user)

    existing_codes = {
        str(doc.get("patrimonial_code", "")).upper()
        for doc in await Equipment.get_pymongo_collection().find(
            {}, {"patrimonial_code": 1}
        ).to_list(None)
    }
    seen_codes: set[str] = set()
    processed = imported = rejected = normalized_rows = 0
    all_errors: list[EquipmentImportError] = []
    all_normalizations: list[EquipmentNormalizationExample] = []

    def value_at(row, key: str):
        idx = headers.get(key)
        return row[idx] if idx is not None and idx < len(row) else None

    try:
        for row_number, row in enumerate(rows, start=2):
            if row_number > 5001:
                all_errors.append(
                    EquipmentImportError(row=row_number, message="Se alcanzó el máximo de 5000 filas por importación.")
                )
                break
            if not any(value not in (None, "") for value in row):
                continue

            processed += 1
            code = _cell_text(value_at(row, "codigo_patrimonial")).upper()
            try:
                if len(code) < 3 or len(code) > 40:
                    raise ValueError("Código patrimonial: debe tener entre 3 y 40 caracteres.")
                if code in seen_codes:
                    raise ValueError("Código patrimonial duplicado dentro del Excel.")
                seen_codes.add(code)
                if code in existing_codes:
                    raise ValueError("El código patrimonial ya existe en el inventario.")

                zone_id_text = _cell_text(value_at(row, "zona_id"))
                zone_name_text = _cell_text(value_at(row, "zona"))
                zone_id = parse_id(zone_id_text) if zone_id_text else None

                office_code_text = _cell_text(value_at(row, "oficina_codigo"))
                office_text = _cell_text(value_at(row, "oficina"))
                office = by_code.get(_normal(office_code_text)) if office_code_text else None
                if not office and office_text:
                    office = by_code.get(_normal(office_text))
                if not office and office_text:
                    matches = by_name.get(_normal(office_text), [])
                    if len(matches) == 1:
                        office = matches[0]
                    elif len(matches) > 1:
                        raise ValueError("El nombre de oficina es ambiguo; use Código Oficina.")
                if not office:
                    raise ValueError("La oficina indicada no existe o está inactiva.")
                if zone_id:
                    if not office.zone_id or office.zone_id != zone_id:
                        raise ValueError("La oficina no pertenece al zona_id indicado en el Excel.")
                elif zone_name_text:
                    if not office.zone_name or _normal(office.zone_name) != _normal(zone_name_text):
                        raise ValueError("La oficina no pertenece a la zona indicada en el Excel.")
                else:
                    raise ValueError("Debe indicar zona o zona_id.")

                ip_value = _cell_text(value_at(row, "ip")) or None
                if ip_value:
                    ipaddress.ip_address(ip_value)

                criticality_raw = value_at(row, "criticidad")
                criticality = 1 if criticality_raw in (None, "") else int(float(criticality_raw))
                if criticality not in (1, 2, 3):
                    raise ValueError("Criticidad: use 1, 2 o 3.")

                responsible_type = (_cell_text(value_at(row, "responsable_tipo")) or "OFICINA").upper()
                responsible_code = _cell_text(value_at(row, "responsable_codigo"))
                responsible_name = _cell_text(value_at(row, "responsable"))
                responsable_id = None
                if responsible_type == "OFICINA":
                    responsible_name = responsible_name or office.name
                elif responsible_type == "JEFE":
                    responsible_name = responsible_name or office.head_name or office.name
                elif responsible_type == "USUARIO":
                    if not responsible_code and not responsible_name:
                        raise ValueError("Responsable: indique el código o nombre del usuario.")
                    office_key = str(office.id)
                    municipal_user = (
                        users_by_office_code.get((office_key, _normal(responsible_code)))
                        if responsible_code
                        else None
                    )
                    if not municipal_user and responsible_name:
                        municipal_user = users_by_office_code.get((office_key, _normal(responsible_name)))
                    if not municipal_user and responsible_name:
                        matches = users_by_office_name.get((office_key, _normal(responsible_name)), [])
                        if len(matches) == 1:
                            municipal_user = matches[0]
                        elif len(matches) > 1:
                            raise ValueError("Responsable ambiguo: use Código Responsable.")
                    if not municipal_user:
                        raise ValueError("El responsable no está registrado como usuario de esa oficina.")
                    responsable_id = municipal_user.id
                    responsible_name = municipal_user.full_name
                else:
                    raise ValueError("responsable_tipo debe ser USUARIO, JEFE u OFICINA.")

                normalization = normalize_inventory_fields(
                    raw_type=_cell_text(value_at(row, "tipo")),
                    device=_cell_text(value_at(row, "dispositivo")),
                    brand=_cell_text(value_at(row, "marca")),
                    model=_cell_text(value_at(row, "modelo")),
                )
                equipment = Equipment(
                    inventory_id=await _next_inventory_id(),
                    patrimonial_code=code,
                    codigo_patrimonial=code,
                    type=normalization.equipment_type,
                    area=_cell_text(value_at(row, "area")) or office.name,
                    device_label=normalization.device_label,
                    brand=normalization.brand,
                    model=_cell_text(value_at(row, "modelo")) or None,
                    hostname=_cell_text(value_at(row, "nombre_equipo")) or _cell_text(value_at(row, "hostname")) or None,
                    ip_address=ip_value,
                    mac_address=_cell_text(value_at(row, "mac")) or None,
                    office_id=office.id,
                    responsable_id=responsable_id,
                    responsible_name=responsible_name,
                    responsible_type=responsible_type,
                    property_type=_cell_text(value_at(row, "propiedad")) or "MUNICIPALIDAD PROVINCIAL DE CASMA",
                    specs={
                        "cpu": _cell_text(value_at(row, "cpu")) or None,
                        "ram_gb": _optional_float(value_at(row, "ram_gb"), "RAM"),
                        "storage_gb": _optional_float(value_at(row, "almacenamiento_gb"), "Almacenamiento"),
                        "screen_size_inches": _optional_float(value_at(row, "tamano_pantalla"), "Tamaño de pantalla"),
                        "os": _cell_text(value_at(row, "sistema_operativo")) or None,
                    },
                    acquired_on=_parse_excel_date(value_at(row, "fecha_adquisicion"), "Fecha de adquisición"),
                    warranty_until=_parse_excel_date(value_at(row, "garantia_hasta"), "Garantía"),
                    status=_equipment_status(value_at(row, "estado")),
                    criticality=criticality,
                    notes=_cell_text(value_at(row, "notas")) or None,
                )
                await equipment.insert()
                if normalization.changes:
                    normalized_rows += 1
                    for change in normalization.changes:
                        all_normalizations.append(
                            EquipmentNormalizationExample(
                                row=row_number,
                                field=change.field,
                                original=change.original,
                                normalized=change.normalized,
                                method=change.method,
                            )
                        )
                existing_codes.add(code)
                imported += 1
            except (ValueError, ValidationError, DuplicateKeyError) as exc:
                rejected += 1
                if isinstance(exc, ValidationError):
                    message = "; ".join(error["msg"].replace("Value error, ", "") for error in exc.errors())
                elif isinstance(exc, DuplicateKeyError):
                    message = "El código patrimonial ya existe en el inventario."
                else:
                    message = str(exc)
                all_errors.append(
                    EquipmentImportError(
                        row=row_number,
                        patrimonial_code=code or None,
                        message=message,
                    )
                )
    finally:
        workbook.close()

    await audit.record(
        request,
        "staff",
        "equipment.imported_xlsx",
        actor_id=str(admin.id),
        actor_name=admin.full_name,
        target_type="equipment",
        target_id="bulk",
        processed=processed,
        imported=imported,
        rejected=rejected,
        normalized=normalized_rows,
        normalization_changes=len(all_normalizations),
        filename=(file.filename or "")[:120],
    )

    visible_errors = all_errors[:200]
    visible_normalizations = all_normalizations[:200]
    return EquipmentImportResult(
        processed=processed,
        imported=imported,
        rejected=rejected,
        normalized=normalized_rows,
        normalizations=visible_normalizations,
        more_normalizations=max(0, len(all_normalizations) - len(visible_normalizations)),
        errors=visible_errors,
        more_errors=max(0, len(all_errors) - len(visible_errors)),
    )


@router.get("/{equipment_id}/retirement-report", response_model=EquipmentRetirementReport)
async def equipment_retirement_report(
    equipment_id: str,
    _: StaffUser = Depends(require_staff),
):
    eq = await _get(equipment_id)
    office = await Office.get(eq.office_id) if eq.office_id else None
    tickets = await Ticket.find(
        {"equipment_id": eq.id, "deleted_at": None}
    ).sort(-Ticket.created_at).to_list()

    now = utcnow()
    incidents_365d = sum(1 for ticket in tickets if ticket.created_at >= now - timedelta(days=365))
    incidents_90d = sum(1 for ticket in tickets if ticket.created_at >= now - timedelta(days=90))
    resolved = [ticket for ticket in tickets if ticket.status.value == "RESUELTO" and ticket.resolution]

    resolution_counts = Counter(
        ticket.resolution.tipo_resolucion.value
        for ticket in resolved
        if ticket.resolution
    )

    engine = get_engine()
    await engine.ensure_ready()
    history_rows = await load_ticket_rows(now - timedelta(days=730), {"equipment_id": eq.id})
    risk, risk_factors = engine.state["risk"].score(
        equipment_row(eq),
        history_rows,
        engine.state["model_rates"],
        now,
    )

    indicators: list[str] = []
    explicit_retirement = (
        resolution_counts.get("OBSOLETO", 0)
        + resolution_counts.get("IRREPARABLE", 0)
        + resolution_counts.get("BAJA_PATRIMONIAL", 0)
    )
    if eq.status.value == "BAJA":
        indicators.append("El activo ya figura con estado BAJA en el inventario TI.")
    if resolution_counts.get("OBSOLETO", 0):
        indicators.append(f"{resolution_counts['OBSOLETO']} incidencia(s) fueron cerradas como equipo obsoleto.")
    if resolution_counts.get("IRREPARABLE", 0):
        indicators.append(f"{resolution_counts['IRREPARABLE']} incidencia(s) fueron cerradas como irreparables.")
    if resolution_counts.get("BAJA_PATRIMONIAL", 0):
        indicators.append(f"{resolution_counts['BAJA_PATRIMONIAL']} incidencia(s) registran baja patrimonial como resultado.")
    if resolution_counts.get("REPARADO", 0) >= 2:
        indicators.append(f"El equipo registra {resolution_counts['REPARADO']} reparaciones documentadas.")
    if resolution_counts.get("REQUIERE_REPUESTO", 0):
        indicators.append(
            f"{resolution_counts['REQUIERE_REPUESTO']} incidencia(s) fueron cerradas indicando necesidad de repuesto."
        )
    if incidents_365d >= 4:
        indicators.append(f"Alta recurrencia: {incidents_365d} incidencias en los últimos 365 días.")
    if risk is not None and risk >= 0.6:
        indicators.append(f"El modelo de riesgo estima {round(risk * 100)}% de probabilidad de nueva falla en 30 días.")

    if explicit_retirement or eq.status.value == "BAJA":
        recommendation = "EVALUAR_BAJA_PATRIMONIAL"
        technical_conclusion = (
            "El historial contiene evidencia técnica que justifica elevar el activo a revisión para baja patrimonial. "
            "La decisión administrativa final corresponde a las áreas competentes."
        )
    elif incidents_365d >= 4 or resolution_counts.get("REPARADO", 0) >= 2 or (risk is not None and risk >= 0.6):
        recommendation = "REQUIERE_EVALUACION_TECNICA"
        technical_conclusion = (
            "El activo presenta recurrencia o riesgo suficiente para una evaluación técnica integral. "
            "Se recomienda documentar diagnóstico, costo de reparación y condición del bien antes de proponer una baja."
        )
    else:
        recommendation = "SIN_EVIDENCIA_SUFICIENTE_PARA_BAJA"
        technical_conclusion = (
            "Con la información registrada actualmente no existe evidencia técnica suficiente para sugerir una baja patrimonial. "
            "Debe continuarse el seguimiento y documentar futuras intervenciones."
        )

    return EquipmentRetirementReport(
        generated_at=now,
        equipment=equipment_out(
            eq,
            office.name if office else None,
            str(office.zone_id) if office and office.zone_id else None,
            office.zone_name if office else None,
        ),
        total_incidents=len(tickets),
        resolved_incidents=len(resolved),
        incidents_365d=incidents_365d,
        incidents_90d=incidents_90d,
        resolution_counts=dict(resolution_counts),
        risk_30d=risk,
        risk_factors=risk_factors,
        indicators=indicators,
        recommendation=recommendation,
        technical_conclusion=technical_conclusion,
        disclaimer=(
            "Este reporte constituye sustento técnico del Área TI y no reemplaza los procedimientos, "
            "informes ni autorizaciones administrativas que correspondan para la disposición o baja de bienes patrimoniales."
        ),
        history=[
            RetirementHistoryItem(
                number=ticket.number,
                created_at=ticket.created_at,
                subject=ticket.subject,
                status=ticket.status.value,
                resolution_type=ticket.resolution.tipo_resolucion.value if ticket.resolution else None,
                resolution_notes=ticket.resolution.notes if ticket.resolution else None,
            )
            for ticket in tickets
        ],
    )


@router.get("/{equipment_id}", response_model=EquipmentDetailOut)
async def equipment_detail(equipment_id: str, _: StaffUser = Depends(require_staff)):
    eq = await _get(equipment_id)
    engine = get_engine()
    await engine.ensure_ready()
    history = await load_ticket_rows(utcnow().replace(year=utcnow().year - 2), {"equipment_id": eq.id})
    risk, factors = engine.state["risk"].score(equipment_row(eq), history, engine.state["model_rates"], utcnow())
    tickets = await Ticket.find({"equipment_id": eq.id, "deleted_at": None}).sort(-Ticket.created_at).limit(50).to_list()
    office = await Office.get(eq.office_id) if eq.office_id else None
    return EquipmentDetailOut(
        equipment=equipment_out(
            eq,
            office.name if office else None,
            str(office.zone_id) if office and office.zone_id else None,
            office.zone_name if office else None,
        ),
        risk=risk,
        risk_factors=factors,
        tickets=[ticket_out(t) for t in tickets],
    )


@router.put("/{equipment_id}", response_model=EquipmentOut)
async def update_equipment(request: Request, equipment_id: str, body: EquipmentIn, user: StaffUser = Depends(require_staff)):
    _ensure_supported_type(body.type)
    eq = await _get(equipment_id)
    office = await _validate_office(body.office_id)
    data = await _assignment_data(body.model_dump(), office)
    for key, value in data.items():
        setattr(eq, key, value)
    eq.updated_at = utcnow()
    try:
        await eq.save()
    except DuplicateKeyError:
        raise HTTPException(status_code=409, detail="Ya existe un equipo con ese código patrimonial.")
    await audit.record(request, "staff", "equipment.updated", actor_id=str(user.id), actor_name=user.full_name, target_type="equipment", target_id=str(eq.id))
    return equipment_out(
        eq,
        office.name,
        str(office.zone_id) if office.zone_id else None,
        office.zone_name,
    )


@router.delete("/{equipment_id}", response_model=Message)
async def delete_equipment(request: Request, equipment_id: str, admin: StaffUser = Depends(require_admin)):
    eq = await _get(equipment_id)
    if await Ticket.find({"equipment_id": eq.id}).count():
        raise HTTPException(status_code=409, detail="El equipo tiene historial de incidencias. Márquelo como BAJA en lugar de eliminarlo.")
    await Device.get_pymongo_collection().update_many({"equipment_id": eq.id}, {"$set": {"equipment_id": None}})
    await eq.delete()
    await audit.record(request, "staff", "equipment.deleted", actor_id=str(admin.id), actor_name=admin.full_name, target_type="equipment", target_id=str(eq.id))
    return Message(message="Equipo eliminado.")


@router.get("/{equipment_id}/qr")
async def equipment_qr(equipment_id: str, _: StaffUser = Depends(require_staff)):
    eq = await _get(equipment_id)
    url = f"{get_settings().public_base_url.rstrip('/')}/reportar?equipo={eq.patrimonial_code}"
    buffer = io.BytesIO()
    qrcode.make(url, box_size=10, border=2).save(buffer, format="PNG")
    return Response(buffer.getvalue(), media_type="image/png", headers={"Content-Disposition": f'inline; filename="QR-{eq.patrimonial_code}.png"'})
